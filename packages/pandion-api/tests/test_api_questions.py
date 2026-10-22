"""A question set kept against its document, checked against the document's text."""

import io
import json

from apihelpers import make_client, upload
from pandion_api import Policy

PASSAGE = "The capital of France is Paris."


def docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Ferns", 1)
    document.add_paragraph("Ferns need water weekly.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def post_set(client, sha, questions, filename="q.json"):
    blob = json.dumps(questions).encode("utf-8")
    return client.post(f"/documents/{sha}/questions", files={"file": (filename, blob, "application/json")})


def test_a_set_is_checked_kept_read_back_and_removed(client, tmp_path):
    sha = upload(client)["sha"]
    r = post_set(client, sha, [{"id": "capital", "question": "Capital?", "gold_answers": [PASSAGE]}])
    body = r.json()
    assert r.status_code == 200 and body["stored"] is True and body["note"] == ""
    assert body["parser"] == "pdfium" and "pdfium" in body["parser_note"]
    assert body["summary"] == {"questions": 1, "found": 1, "found_normalized": 0, "not_found": 0}
    assert (tmp_path / "data" / "sources" / "questions" / f"{sha}.json").is_file()
    assert client.get(f"/documents/{sha}/questions").json()["count"] == 1
    assert client.delete(f"/documents/{sha}/questions").status_code == 204
    assert client.get(f"/documents/{sha}/questions").json() is None
    assert client.delete(f"/documents/{sha}/questions").status_code == 204


def test_a_passage_that_is_not_in_the_document_is_reported_and_the_set_still_kept(client):
    sha = upload(client)["sha"]
    body = post_set(client, sha, [{"question": "Q?", "gold_answers": ["Nothing like this is in it."]}]).json()
    assert body["summary"]["not_found"] == 1 and body["stored"] is True


def test_a_bad_set_is_400_with_its_sentence(client):
    sha = upload(client)["sha"]
    r = post_set(client, sha, [{"question": "", "gold_answers": [PASSAGE]}])
    assert r.status_code == 400 and r.json()["detail"]


def test_a_text_document_is_read_by_docling_and_checked(client):
    sha = client.post("/documents", files={"file": ("a.txt", b"hello there", "text/plain")}).json()["sha"]
    body = post_set(client, sha, [{"question": "Q?", "gold_answers": ["hello there"]}]).json()
    assert body["summary"]["found"] == 1
    assert body["parser"] == "docling" and "docling parser" in body["parser_note"]


def test_a_word_document_is_checked_as_docling_reads_it(client):
    sha = client.post("/documents", files={"file": ("field-guide.docx", docx_bytes(), "application/octet-stream")}).json()["sha"]
    body = post_set(client, sha, [{"question": "Water?", "gold_answers": ["Ferns need water weekly."]}]).json()
    assert body["summary"] == {"questions": 1, "found": 1, "found_normalized": 0, "not_found": 0}
    assert body["parser"] == "docling"


def test_an_image_is_415_because_reading_it_needs_ocr_models(client):
    sha = client.post("/documents", files={"file": ("scan.png", b"PNG not really", "image/png")}).json()["sha"]
    r = post_set(client, sha, [{"question": "Q?", "gold_answers": ["hello there"]}])
    assert r.status_code == 415 and "image" in r.json()["detail"]


def test_a_broken_word_file_is_415_because_it_cannot_be_read(client):
    sha = client.post(
        "/documents", files={"file": ("notes.docx", b"PK not really a word file", "application/octet-stream")}
    ).json()["sha"]
    r = post_set(client, sha, [{"question": "Q?", "gold_answers": ["hello there"]}])
    assert r.status_code == 415


def test_an_unknown_or_hidden_document_is_404(tmp_path):
    class Hidden(Policy):
        def can_read(self, sha, request):
            return False

    with make_client(tmp_path, Hidden()) as c:
        sha = upload(c)["sha"]
        assert c.get(f"/documents/{sha}/questions").status_code == 404
        assert post_set(c, sha, [{"question": "Q?", "gold_answers": [PASSAGE]}]).status_code == 404
        assert c.get(f"/documents/{'0' * 64}/questions").status_code == 404


def test_a_policy_can_check_a_set_without_keeping_it(tmp_path):
    class Keepless(Policy):
        not_keeping_note = "Checked, not kept."

        def keeps_question_sets(self, request):
            return False

    with make_client(tmp_path, Keepless()) as c:
        sha = upload(c)["sha"]
        body = post_set(c, sha, [{"question": "Capital?", "gold_answers": [PASSAGE]}]).json()
        assert body["stored"] is False and body["note"] == "Checked, not kept."
        assert body["summary"]["found"] == 1
        assert c.get(f"/documents/{sha}/questions").json() is None
    assert not (tmp_path / "data" / "sources" / "questions").exists()


def test_find_gold_places_a_passage_by_the_eval_steps_rule(client):
    pieces = ["Nothing here.", "the capital of france is “paris.”", PASSAGE]
    r = client.post("/golds/find", json={"golds": ['The capital of France is "Paris."'], "texts": pieces})
    assert r.status_code == 200
    assert r.json() == {"found": {"index": 1, "gold": 'The capital of France is "Paris."', "match": "normalized"}}


def test_find_gold_says_null_when_no_piece_holds_one(client):
    assert client.post("/golds/find", json={"golds": [PASSAGE], "texts": ["Lyon."]}).json() == {"found": None}
    assert client.post("/golds/find", json={"golds": [PASSAGE], "texts": []}).json() == {"found": None}


def test_find_gold_refuses_what_it_does_not_look_through(client):
    assert client.post("/golds/find", json={"golds": [], "texts": ["x"]}).status_code == 422
    assert client.post("/golds/find", json={"golds": ["x"], "texts": ["y"] * 101}).status_code == 422
    assert client.post("/golds/find", json={"golds": ["x" * 10_001], "texts": []}).status_code == 422
    assert client.post("/golds/find", json={"golds": ["x"], "texts": ["y" * 500_000, "y" * 500_001]}).status_code == 422


def test_find_gold_caps_keep_one_call_small(client):
    """The caps bound one call's work: 20 passages of at most 10,000 characters,
    and 100 pieces of at most 1,000,000 characters together. One long piece
    is looked through whole."""
    from pandion_api.routes.questions import MAX_GOLDS, MAX_PIECES, MAX_TEXT_CHARS, MAX_TOTAL_CHARS

    assert (MAX_GOLDS, MAX_PIECES, MAX_TEXT_CHARS, MAX_TOTAL_CHARS) == (20, 100, 10_000, 1_000_000)
    body = {"golds": ["g" * MAX_TEXT_CHARS] * MAX_GOLDS, "texts": ["t" * 10_000] * MAX_PIECES}
    assert client.post("/golds/find", json=body).json() == {"found": None}
    long = "t" * 900_000 + " the gold passage"
    found = client.post("/golds/find", json={"golds": ["the gold passage"], "texts": [long]}).json()["found"]
    assert found == {"index": 0, "gold": "the gold passage", "match": "exact"}


# --- one set for several documents (multi-document part 4) ------------------------

from pdfs import build_pdf  # noqa: E402

BEES = "Bees dance to show the hive where the flowers are."


def two(client):
    first = upload(client)["sha"]
    second = upload(client, "bees.pdf", build_pdf([[BEES], ["Honey keeps for years."]]))["sha"]
    return first, second


def post_across(client, shas, questions, filename="golden.json"):
    blob = json.dumps(questions).encode("utf-8")
    return client.post(
        "/questions",
        data={"documents": list(shas)},
        files={"file": (filename, blob, "application/json")},
    )


def test_one_set_for_several_documents_is_checked_each_question_in_its_own_document(client):
    first, second = two(client)
    body = post_across(client, [first, second], [
        {"id": "capital", "question": "Capital?", "gold_answers": [PASSAGE], "document": "report.pdf"},
        {"id": "wrong", "question": "Bees?", "gold_answers": [BEES], "document": "report.pdf"},
        {"id": "bees", "question": "Bees?", "gold_answers": [BEES], "document": second},
        {"id": "any", "question": "Bees anywhere?", "gold_answers": [BEES]},
    ]).json()
    assert body["stored"] is True and body["documents"] == [first, second]
    assert body["summary"] == {"questions": 4, "found": 3, "found_normalized": 0, "not_found": 1}
    rows = {r["id"]: r for r in body["questions"]}
    assert (rows["capital"]["document"], rows["capital"]["golds"][0]["document"]) == (first, first)
    assert (rows["wrong"]["status"], rows["wrong"]["golds"][0]["document"]) == ("not_found", None)
    assert (rows["bees"]["document"], rows["bees"]["golds"][0]["document"]) == (second, second)
    assert (rows["any"]["document"], rows["any"]["golds"][0]["document"]) == (None, second)


def test_one_set_is_kept_split_by_document_with_unnamed_questions_in_each(client):
    first, second = two(client)
    post_across(client, [first, second], [
        {"id": "capital", "question": "Capital?", "gold_answers": [PASSAGE], "document": "report.pdf"},
        {"id": "bees", "question": "Bees?", "gold_answers": [BEES], "document": "bees.pdf"},
        {"id": "any", "question": "Bees anywhere?", "gold_answers": [BEES]},
    ])
    kept = {sha: client.get(f"/documents/{sha}/questions").json() for sha in (first, second)}
    assert [(q["id"], q["document"]) for q in kept[first]["set"]["questions"]] == [("capital", "report.pdf"), ("any", "")]
    assert [(q["id"], q["document"]) for q in kept[second]["set"]["questions"]] == [("bees", "bees.pdf"), ("any", "")]
    assert kept[first]["filename"] == kept[second]["filename"] == "golden.json"


def test_a_document_the_new_set_has_no_questions_for_loses_its_old_set(client):
    first, second = two(client)
    post_set(client, second, [{"id": "old", "question": "Old?", "gold_answers": [BEES]}])
    post_across(client, [first, second], [
        {"id": "capital", "question": "Capital?", "gold_answers": [PASSAGE], "document": "report.pdf"},
    ])
    assert client.get(f"/documents/{second}/questions").json() is None


def test_a_set_naming_a_document_it_is_not_for_is_400_and_nothing_is_kept(client):
    first, second = two(client)
    r = post_across(client, [first, second], [
        {"question": "Capital?", "gold_answers": [PASSAGE], "document": "elsewhere.pdf"},
    ])
    assert r.status_code == 400
    assert r.json()["detail"] == (
        "Question 1 names the document elsewhere.pdf, which is not one of the documents "
        "searched: bees.pdf, report.pdf."
    )
    assert client.get(f"/documents/{first}/questions").json() is None


def test_one_set_across_a_pdf_and_a_markdown_file_is_checked_in_each(client):
    pdf = upload(client)["sha"]
    md = client.post("/documents", files={"file": ("notes.md", f"# Notes\n\n{BEES}\n".encode(), "text/markdown")}).json()["sha"]
    body = post_across(client, [pdf, md], [
        {"id": "capital", "question": "Capital?", "gold_answers": [PASSAGE], "document": "report.pdf"},
        {"id": "bees", "question": "Bees?", "gold_answers": [BEES], "document": "notes.md"},
    ]).json()
    assert body["summary"]["found"] == 2
    assert [r["golds"][0]["document"] for r in body["questions"]] == [pdf, md]
    assert body["parser"] == "docling, pdfium" and "pdfium for a PDF, docling for the others" in body["parser_note"]


def test_a_pdf_named_as_markdown_is_415_and_never_reaches_docling(client, monkeypatch):
    import pandion

    seen = []
    monkeypatch.setattr(pandion, "parse", lambda *a, **k: seen.append(k) or (_ for _ in ()).throw(AssertionError))
    sha = client.post("/documents", files={"file": ("notes.md", build_pdf([[PASSAGE]]), "text/markdown")}).json()["sha"]
    r = post_set(client, sha, [{"question": "Q?", "gold_answers": [PASSAGE]}])
    assert r.status_code == 415
    assert r.json()["detail"] == "this file is named .md but holds a PDF, so its text cannot be checked"
    assert seen == []


def test_an_image_named_as_text_is_415(client):
    png = b"\x89PNG\r\n\x1a\n" + bytes(64)
    sha = client.post("/documents", files={"file": ("notes.txt", png, "text/plain")}).json()["sha"]
    r = post_set(client, sha, [{"question": "Q?", "gold_answers": ["hello"]}])
    assert r.status_code == 415
    assert r.json()["detail"] == "this file is named .txt but holds an image, so its text cannot be checked"


def test_the_set_for_several_documents_comes_back_as_one(client):
    first, second = two(client)
    post_across(client, [first, second], [
        {"id": "capital", "question": "Capital?", "gold_answers": [PASSAGE], "document": "report.pdf"},
        {"id": "bees", "question": "Bees?", "gold_answers": [BEES], "document": "bees.pdf"},
        {"id": "any", "question": "Bees anywhere?", "gold_answers": [BEES]},
    ])
    body = client.get("/questions", params={"documents": [first, second]}).json()
    assert body["documents"] == [first, second] and body["count"] == 3
    assert [(q["id"], q["document"]) for q in body["questions"]] == [("capital", first), ("any", ""), ("bees", second)]


def test_unnamed_questions_stay_when_the_first_document_goes(client):
    first, second = two(client)
    post_across(client, [first, second], [{"id": "any", "question": "Bees anywhere?", "gold_answers": [BEES]}])
    store = client.app.state.documents
    store.remove(first)
    body = client.get("/questions", params={"documents": [second]}).json()
    assert [q["id"] for q in body["questions"]] == ["any"]
